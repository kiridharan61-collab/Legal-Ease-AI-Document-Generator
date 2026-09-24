from io import BytesIO
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Inches, Pt
from fpdf import FPDF

from utils.sanitize import sanitize_text


def _clean_filename(value: str) -> str:
    value = re.sub(r"[^A-Za-z0-9_-]+", "_", value.strip())
    return value.strip("_") or "legalease_document"


def format_txt(text: str) -> bytes:
    return sanitize_text(text).encode("utf-8")


def _add_docx_header(document: Document, doc_type: str, logo_bytes: bytes | None):
    section = document.sections[0]
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if logo_bytes:
        run = paragraph.add_run()
        run.add_picture(BytesIO(logo_bytes), width=Inches(1.0))

    title = paragraph.add_run(f"\n{sanitize_text(doc_type).upper()}")
    title.bold = True
    title.font.name = "Times New Roman"
    title.font.size = Pt(14)


def _add_docx_footer(document: Document):
    for section in document.sections:
        paragraph = section.footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run("Generated with LegalEase | Review before signing")
        run.font.name = "Times New Roman"
        run.font.size = Pt(8)


def format_docx(
    text: str,
    doc_type: str,
    terms: str = "",
    logo_bytes: bytes | None = None,
) -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = document.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(11)

    _add_docx_header(document, doc_type, logo_bytes)

    for raw_line in sanitize_text(text).splitlines():
        line = raw_line.strip()
        if not line:
            document.add_paragraph("")
            continue

        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(6)

        if (
            line.isupper()
            or re.match(r"^\d+[\.\)]\s+", line)
            or line.lower().startswith(("review notice", "signature"))
        ):
            run = p.add_run(line)
            run.bold = True
        elif line.startswith("- "):
            p.style = document.styles["List Bullet"]
            p.add_run(line[2:].strip())
        else:
            p.add_run(line)

        for run in p.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(11)

    clean_terms = [t for t in terms.split(";") if t.strip()]
    if clean_terms:
        document.add_paragraph("")
        heading = document.add_paragraph()
        heading.add_run("KEY TERMS").bold = True

        table = document.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"
        table.rows[0].cells[0].text = "No."
        table.rows[0].cells[1].text = "Term / Condition"

        for index, term in enumerate(clean_terms, 1):
            cells = table.add_row().cells
            cells[0].text = str(index)
            cells[1].text = sanitize_text(term)

    _add_docx_footer(document)

    output = BytesIO()
    document.save(output)
    return output.getvalue()


class LegalEasePDF(FPDF):
    def __init__(self, doc_type: str):
        super().__init__()
        self.doc_type = sanitize_text(doc_type)
        self.set_auto_page_break(auto=True, margin=18)

    def header(self):
        self.set_font("Helvetica", "B", 12)
        self.cell(0, 8, self.doc_type.upper(), align="C")
        self.ln(10)

    def footer(self):
        self.set_y(-14)
        self.set_font("Helvetica", "", 8)
        self.cell(
            0,
            8,
            f"LegalEase | Page {self.page_no()} | Review before signing",
            align="C",
        )


def format_pdf(
    text: str,
    doc_type: str,
    logo_path: str | None = None,
) -> bytes:
    pdf = LegalEasePDF(doc_type)
    pdf.add_page()

    if logo_path and Path(logo_path).exists():
        try:
            pdf.image(logo_path, x=pdf.w / 2 - 12, y=18, w=24)
            pdf.ln(22)
        except Exception:
            pass

    for raw_line in sanitize_text(text).splitlines():
        line = raw_line.strip()
        if not line:
            pdf.ln(3)
            continue

        is_heading = (
            line.isupper()
            or bool(re.match(r"^\d+[\.\)]\s+", line))
            or line.lower().startswith(("review notice", "signature"))
        )

        pdf.set_font("Helvetica", "B" if is_heading else "", 11)
        if line.startswith("- "):
            pdf.multi_cell(0, 6, "- " + line[2:].strip())
        else:
            pdf.multi_cell(0, 6, line)
        pdf.ln(1)

    return bytes(pdf.output())
